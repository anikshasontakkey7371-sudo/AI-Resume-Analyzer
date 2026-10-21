"""
Generates sample PDF and DOCX resumes in sample_data directory.
"""
import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document

from sample_data import SAMPLE_RESUMES

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "sample_data")
os.makedirs(OUTPUT_DIR, exist_ok=True)

def generate_pdf():
    pdf_path = os.path.join(OUTPUT_DIR, "sample_resume.pdf")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'ResumeTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=4
    )
    
    heading_style = ParagraphStyle(
        'ResumeHeading',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#4338ca"),
        spaceBefore=8,
        spaceAfter=4
    )
    
    body_style = ParagraphStyle(
        'ResumeBody',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#1e293b")
    )
    
    elements = []
    resume_text = SAMPLE_RESUMES["software_engineer"]["text"]
    
    lines = resume_text.split('\n')
    for line in lines:
        line_s = line.strip()
        if not line_s:
            elements.append(Spacer(1, 4))
        elif line_s.isupper() and len(line_s) > 3:
            elements.append(Paragraph(line_s, heading_style))
        elif "@" in line_s or "linkedin.com" in line_s:
            elements.append(Paragraph(line_s, body_style))
        elif line_s == "Alex Morgan":
            elements.append(Paragraph(line_s, title_style))
        else:
            elements.append(Paragraph(line_s, body_style))
            
    doc.build(elements)
    print(f"Generated PDF: {pdf_path}")

def generate_docx():
    docx_path = os.path.join(OUTPUT_DIR, "sample_resume.docx")
    doc = Document()
    resume_text = SAMPLE_RESUMES["data_scientist"]["text"]
    
    lines = resume_text.split('\n')
    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if line_s == "Samantha Chen":
            doc.add_heading(line_s, level=0)
        elif line_s.isupper() and len(line_s) > 3:
            doc.add_heading(line_s, level=1)
        elif line_s.startswith("- "):
            doc.add_paragraph(line_s[2:], style='List Bullet')
        else:
            doc.add_paragraph(line_s)
            
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")

if __name__ == "__main__":
    generate_pdf()
    generate_docx()
