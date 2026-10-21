"""
Resume Parser Service.
Extracts readable text from PDF and DOCX files with sanitization and validation.
"""
import io
import os
import sys
from typing import Dict, Any, Optional

_BACKEND_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if _BACKEND_DIR not in sys.path:
    sys.path.insert(0, _BACKEND_DIR)

# Resilient PDF parser imports
try:
    from pypdf import PdfReader
except ImportError:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        PdfReader = None

# Resilient DOCX parser imports
try:
    from docx import Document
except ImportError:
    Document = None

try:
    from utils.validators import validate_file, FileValidationError
except ImportError:
    from backend.utils.validators import validate_file, FileValidationError

class ParsingError(Exception):
    """Raised when document parsing fails."""
    pass

class ResumeParser:
    """Service to handle PDF and DOCX text extraction."""

    @staticmethod
    def extract_from_pdf(file_stream) -> str:
        """Extract plain text from a PDF stream using pypdf/PyPDF2."""
        if PdfReader is None:
            raise ParsingError("PDF parser library (pypdf/PyPDF2) is not available.")
            
        try:
            if hasattr(file_stream, 'seek'):
                file_stream.seek(0)
                
            reader = PdfReader(file_stream)
            text_parts = []
            
            for page in reader.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_parts.append(page_text.strip())
                    
            full_text = "\n\n".join(text_parts).strip()
            if not full_text:
                raise ParsingError(
                    "Unable to extract readable text from PDF. The document may be scanned, image-only, or encrypted."
                )
            return full_text
        except ParsingError:
            raise
        except Exception as e:
            raise ParsingError(f"Failed to read PDF document: {str(e)}")

    @staticmethod
    def extract_from_docx(file_stream) -> str:
        """Extract plain text from a DOCX stream including paragraphs and tables."""
        if Document is None:
            raise ParsingError("Word document parser library (python-docx) is not available.")
            
        try:
            if hasattr(file_stream, 'seek'):
                file_stream.seek(0)
                
            doc = Document(file_stream)
            text_parts = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    text_parts.append(text)
                    
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)
                        
            full_text = "\n".join(text_parts).strip()
            if not full_text:
                raise ParsingError("The Word document (.docx) does not contain any readable text.")
            return full_text
        except ParsingError:
            raise
        except Exception as e:
            raise ParsingError(f"Failed to read Word document (.docx): {str(e)}")

    @classmethod
    def extract_from_storage(cls, file_storage) -> Dict[str, Any]:
        """
        Validate and extract text from a Flask FileStorage object.
        Returns dictionary with text, statistics, and metadata.
        """
        filename = getattr(file_storage, 'filename', None) or "document"
        ext = os.path.splitext(filename)[1].lower()
        
        file_bytes = file_storage.read()
        file_size = len(file_bytes)
        
        if hasattr(file_storage, 'seek'):
            file_storage.seek(0)
            
        is_valid, err_msg = validate_file(filename, file_size)
        if not is_valid:
            raise FileValidationError(err_msg)
            
        stream = io.BytesIO(file_bytes)
        
        if ext == '.pdf':
            text = cls.extract_from_pdf(stream)
        elif ext == '.docx':
            text = cls.extract_from_docx(stream)
        else:
            raise FileValidationError(f"Unsupported file format '{ext}'.")
            
        words = text.split()
        return {
            "filename": filename,
            "extension": ext,
            "file_size_bytes": file_size,
            "text": text,
            "word_count": len(words),
            "char_count": len(text),
            "status": "success"
        }
